# -*- coding: utf-8 -*-
"""
GENERA LAS ESPECIFICACIONES DE LAS FICHAS DESDE EL EXCEL DE DATOS TECNICOS
===========================================================================
Lee  _tools/especificaciones/datos-tecnicos.xlsx  (hoja "Todas": una fila por
codigo) y escribe  ruta-productos/JS/especificaciones.json  con las 5 medidas
que muestran las fichas de producto:

    ext     = Ø D (mm)                    -> fila "Ø exterior"
    int     = d agujero / Ø interior (mm) -> fila "Ø interior"
    ancho   = B ancho de corte / # (mm)   -> fila "Ancho de corte"
    dientes = Z dientes / filos           -> fila "Cantidad de dientes"
    largo   = Largo L / LT (mm)           -> fila "Largo"

producto.js lee ese JSON y, para el codigo de cada ficha (y cada variante del
desplegable), usa estas medidas como fuente de verdad. Si un codigo no esta en
el Excel, la ficha sigue mostrando lo de antes (no se rompe nada).

El Excel es la UNICA fuente: cuando llega una lista nueva, se reemplaza
datos-tecnicos.xlsx y se vuelve a correr este script (o el .cmd de al lado).

USO:   python generar-especificaciones.py
       (o doble clic en generar-especificaciones.cmd)

Requisitos: Python 3 + openpyxl  (pip install openpyxl)
"""
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instalalo con:  pip install openpyxl")

AQUI = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(AQUI, "datos-tecnicos.xlsx")
SALIDA = os.path.normpath(os.path.join(AQUI, "..", "..", "ruta-productos", "JS", "especificaciones.json"))

# Nombre de hoja y columnas por su encabezado (se buscan por un fragmento
# distintivo para no depender del orden exacto de columnas).
HOJA = "Todas"
MAPEO = {
    "ext":     "D (mm)",          # "Ø D (mm)"
    "ancho":   "B ancho",         # "B ancho de corte / # (mm)"
    "int":     "agujero",         # "d agujero / Ø interior (mm)"
    "dientes": "Z dientes",       # "Z dientes / filos"
    "largo":   "Largo L",         # "Largo L / LT (mm)"
}
# Algunas mechas (MDD, MID...) no cargan "Largo L / LT" sino "LU largo útil":
# si "largo" queda vacío, se toma de ahí para no perder esa medida.
COL_LARGO_ALT = "LU largo"       # "LU largo útil (mm)"
COL_CODIGO = "digo"              # "Código"


def norm(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def main():
    if not os.path.exists(XLSX):
        sys.exit("No encuentro el Excel: " + XLSX +
                 "\nCopia ahi la lista de datos tecnicos como 'datos-tecnicos.xlsx'.")

    wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    if HOJA not in wb.sheetnames:
        sys.exit("El Excel no tiene la hoja '%s'. Hojas: %s" % (HOJA, wb.sheetnames))
    ws = wb[HOJA]

    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        sys.exit("La hoja '%s' esta vacia." % HOJA)
    encabezados = [norm(h) or "" for h in filas[0]]

    def col_idx(fragmento):
        for i, h in enumerate(encabezados):
            if fragmento in h:
                return i
        sys.exit("No encuentro la columna que contenga '%s'. Encabezados: %s" % (fragmento, encabezados))

    idx_cod = col_idx(COL_CODIGO)
    idx = {campo: col_idx(frag) for campo, frag in MAPEO.items()}
    idx_largo_alt = col_idx(COL_LARGO_ALT)

    datos = {}
    duplicados = []
    sin_medida = 0
    for fila in filas[1:]:
        codigo = norm(fila[idx_cod] if idx_cod < len(fila) else None)
        if not codigo:
            continue
        medidas = {}
        for campo, i in idx.items():
            val = norm(fila[i]) if i < len(fila) else None
            if val is not None:
                medidas[campo] = val
        if "largo" not in medidas:
            alt = norm(fila[idx_largo_alt]) if idx_largo_alt < len(fila) else None
            if alt is not None:
                medidas["largo"] = alt
        if not medidas:
            sin_medida += 1
            continue  # sin ninguna medida: la ficha sigue con lo de antes
        if codigo in datos:
            duplicados.append(codigo)
        datos[codigo] = medidas

    # Deteccion de colisiones por codigo compacto (sin espacios, mayusculas):
    # producto.js hace el match asi, por eso avisamos si dos codigos chocan.
    compactos = {}
    colisiones = []
    for c in datos:
        k = c.replace(" ", "").upper()
        if k in compactos and compactos[k] != c:
            colisiones.append((compactos[k], c))
        else:
            compactos[k] = c

    salida = {
        "_meta": {
            "generado_desde": "datos-tecnicos.xlsx (hoja Todas)",
            "campos": "ext=Ø D | int=Ø interior | ancho=B/# | dientes=Z | largo=L/LT",
            "codigos": len(datos),
        },
        "medidas": datos,
    }
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    with open(SALIDA, "w", encoding="utf-8") as f:
        json.dump(salida, f, ensure_ascii=False, separators=(",", ":"), sort_keys=True)

    print("OK  ->", SALIDA)
    print("Codigos con medidas :", len(datos))
    print("Filas sin ninguna medida (se omiten):", sin_medida)
    if duplicados:
        print("Codigos repetidos (se piso con el ultimo):", len(duplicados), duplicados[:10])
    if colisiones:
        print("AVISO colisiones por codigo compacto:", colisiones[:10])


if __name__ == "__main__":
    main()
